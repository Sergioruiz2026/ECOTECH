import os
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from informes.informe import Informe

class InformeExcel(Informe):
    """
    Genera libros de trabajo en formato Microsoft Excel (.xlsx) con estilos corporativos.
    """
    def generar(self, ruta_archivo: str) -> bool:
        try:
            directorio = os.path.dirname(ruta_archivo)
            if directorio and not os.path.exists(directorio):
                os.makedirs(directorio)

            wb = Workbook()
            ws = wb.active
            ws.title = "Reporte ECOTECH"

            # Estilos corporativos
            fuente_titulo = Font(name='Calibri', size=16, bold=True, color='1E3A8A')
            fuente_meta = Font(name='Calibri', size=10, italic=True, color='4B5563')
            fuente_encabezado = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
            fill_encabezado = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
            borde_delgado = Border(
                left=Side(style='thin', color='D1D5DB'),
                right=Side(style='thin', color='D1D5DB'),
                top=Side(style='thin', color='D1D5DB'),
                bottom=Side(style='thin', color='D1D5DB')
            )

            # Título del informe
            ws['A1'] = "ECOTECH Solutions"
            ws['A1'].font = Font(name='Calibri', size=10, bold=True, color='6B7280')
            ws['A2'] = self.titulo
            ws['A2'].font = fuente_titulo
            ws['A3'] = f"Fecha de emisión: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
            ws['A3'].font = fuente_meta

            fila_inicio = 5

            if self.datos:
                encabezados = list(self.datos[0].keys())

                # Escribir encabezados
                for col_idx, encabezado in enumerate(encabezados, start=1):
                    celda = ws.cell(row=fila_inicio, column=col_idx, value=str(encabezado).upper())
                    celda.font = fuente_encabezado
                    celda.fill = fill_encabezado
                    celda.alignment = Alignment(horizontal='center', vertical='center')

                # Escribir registros
                for fila_idx, item in enumerate(self.datos, start=fila_inicio + 1):
                    for col_idx, encabezado in enumerate(encabezados, start=1):
                        valor = item.get(encabezado, '')
                        celda = ws.cell(row=fila_idx, column=col_idx, value=valor)
                        celda.border = borde_delgado
                        
                        # Formato condicional básico para números
                        if isinstance(valor, (int, float)):
                            celda.alignment = Alignment(horizontal='right')
                        else:
                            celda.alignment = Alignment(horizontal='left')

                # Ajustar el ancho automático de las columnas
                for col in ws.columns:
                    max_len = 0
                    col_letter = get_column_letter(col[0].column)
                    for cell in col:
                        if cell.row >= fila_inicio and cell.value:
                            max_len = max(max_len, len(str(cell.value)))
                    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

            wb.save(ruta_archivo)
            print(f"Informe Excel generado exitosamente en: {ruta_archivo}")
            return True

        except Exception as e:
            print(f"Error al generar el informe Excel: {e}")
            return False